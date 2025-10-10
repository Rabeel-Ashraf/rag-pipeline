import os
from typing import Optional
import PyPDF2
from docx import Document
from pptx import Presentation

def extract_text_from_txt(file_path: str) -> str:
    """Extract text from TXT files"""
    with open(file_path, 'r', encoding='utf-8') as file:
        return file.read()

def extract_text_from_pdf(file_path: str) -> str:
    """Extract text from PDF files"""
    text = ""
    try:
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text += extracted + "\n"
    except Exception as e:
        print(f"Error extracting PDF {file_path}: {e}")
    return text

def extract_text_from_docx(file_path: str) -> str:
    """Extract text from DOCX files"""
    text = ""
    try:
        doc = Document(file_path)
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text += paragraph.text + "\n"
    except Exception as e:
        print(f"Error extracting DOCX {file_path}: {e}")
    return text

def extract_text_from_pptx(file_path: str) -> str:
    """Extract text from PPTX files"""
    text = ""
    try:
        prs = Presentation(file_path)
        for slide in prs.slides:
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    text += shape.text + "\n"
    except Exception as e:
        print(f"Error extracting PPTX {file_path}: {e}")
    return text

def extract_text(file_path: str) -> Optional[str]:
    """Extract text based on file extension"""
    _, ext = os.path.splitext(file_path.lower())
    
    if ext == '.txt':
        return extract_text_from_txt(file_path)
    elif ext == '.pdf':
        return extract_text_from_pdf(file_path)
    elif ext == '.docx':
        return extract_text_from_docx(file_path)
    elif ext == '.pptx':
        return extract_text_from_pptx(file_path)
    else:
        return None
